"""流式写入器（源码/xlsx_writer.py）的边界测试。

这个模块直接决定用户拿到的 Excel 能不能打开，而它是新写的、替换掉了原来的第三方库。
最容易出问题的地方是**转义**：工作表名写进 XML 属性、单元格文本写进元素内容，两类上下文
对特殊字符的要求不同。这里逐个钉子。

已修过的一个真实缺陷：店铺名里只要有一个双引号，属性会被提前闭合，整个工作簿损坏打不开——
而 Excel 允许店铺名带双引号（clean_sheet_name 只处理 : \\ / ? * [ ]）。
"""
import json, os, sys, tempfile, unittest
from datetime import datetime
from pathlib import Path

sys.path.insert(0, os.environ.get('COMMERCE_SOURCE', str(Path(__file__).parent.parent / '源码')))
from tables import read_xlsx, sheet_names
import openpyxl
import xlsx_writer


def build(root, sheet_name, headers, rows):
    """写一份最小工作簿，返回路径。"""
    folder = Path(root)
    rows_path = folder / 'rows.jsonl'
    rows_path.write_text(''.join(json.dumps(row, ensure_ascii=False) + '\n' for row in rows), encoding='utf8')
    index = {'headers': headers, 'dateColumns': [], 'numericColumns': [], 'rows': len(rows), 'orders': len(rows),
             'shops': [{'id': '1', 'name': sheet_name, 'rows': len(rows), 'path': str(rows_path),
                        'sheetName': sheet_name}]}
    index_path = folder / 'sheets.json'
    index_path.write_text(json.dumps(index, ensure_ascii=False), encoding='utf8')
    out = folder / 'out.xlsx'
    xlsx_writer.write(out, index_path)
    return out


class Writer(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()

    def tearDown(self):
        self.folder.cleanup()

    def test_sheet_name_with_xml_special_characters(self):
        """店铺名里的双引号/&/尖括号不能让工作簿损坏——它们会进 XML 属性。"""
        name = 'A"B店 & <测试> 再"引号'
        out = build(self.folder.name, name, ['列一'], [['值']])
        book = openpyxl.load_workbook(out)
        self.assertEqual(book.sheetnames, [name], '打开后工作表名必须原样')
        self.assertEqual(sheet_names(out), [name], '项目自己的读取器也要读得出')

    def test_cell_values_with_xml_special_characters(self):
        headers = ['文本', '带引号', '带尖括号与和号']
        values = ['a&b', '他说"你好"', 'x<y>z & w']
        out = build(self.folder.name, '店', headers, [values])
        got = next(iter(openpyxl.load_workbook(out).worksheets)).iter_rows(min_row=2, values_only=True)
        self.assertEqual(list(got)[0], tuple(values))

    def test_leading_and_trailing_spaces_are_kept(self):
        """平台明细里有带前导/尾随空格的取值，不能因为 XML 的空白规则被吃掉。

        两层都要验：一是往返读回不变（openpyxl 与项目读取器本来就不裁剪空白，所以这一层
        测不出 xml:space 的作用）；二是 XML 里确实声明了空格保留——那是给 Excel 这类
        解析器看的，缺了它 Excel 可能把前后空格去掉。"""
        import zipfile
        out = build(self.folder.name, '店', ['列'], [['  前导与尾随  ']])
        heads, rows = read_xlsx(out)
        self.assertEqual(rows[0][0], '  前导与尾随  ')
        with zipfile.ZipFile(out) as archive:
            sheet_xml = archive.read('xl/worksheets/sheet1.xml').decode('utf8')
        self.assertIn('xml:space="preserve"', sheet_xml, '文本单元格要声明保留空格')
        # 表头里的全角空格同理（真实数据里就有）
        out2 = build(self.folder.name, '店', ['链接\u3000'], [['x']])
        self.assertEqual(read_xlsx(out2)[0], ['链接\u3000'])

    def test_control_characters_do_not_break_the_file(self):
        """XML 不允许的控制字符要剥掉：宁可少一个脏字符，也不能产出打不开的工作簿。"""
        out = build(self.folder.name, '店', ['列'], [['前\x00中\x1f后']])
        book = openpyxl.load_workbook(out)
        value = next(iter(book.worksheets)).cell(row=2, column=1).value
        self.assertEqual(value, '前中后')

    def test_value_types_map_to_expected_cell_kinds(self):
        headers = ['订单号', '金额', '发货时间', '空值', '真假']
        values = ['1234567890123456789', 39.9, {'date': '2026-08-31T23:59:59Z'}, '', True]
        out = build(self.folder.name, '店', headers, [values])
        sheet = next(iter(openpyxl.load_workbook(out).worksheets))
        cells = [sheet.cell(row=2, column=i + 1) for i in range(len(headers))]
        self.assertEqual(cells[0].value, '1234567890123456789')
        self.assertEqual(cells[0].data_type, 's', '19 位订单号必须是文本，不能变成科学计数法')
        self.assertEqual(cells[1].value, 39.9)
        self.assertEqual(cells[1].data_type, 'n')
        self.assertEqual(cells[2].value, datetime(2026, 8, 31, 23, 59, 59), '日期必须是真日期，且秒不差')
        self.assertIn(cells[3].value, (None, ''))
        self.assertIs(cells[4].value, True)

    def test_header_row_is_frozen_at_d2(self):
        """冻结窗格在 D2：这是既有测试钉住过的形态，换实现后必须保持。"""
        out = build(self.folder.name, '店', ['a', 'b', 'c', 'd'], [['1', '2', '3', '4']])
        self.assertEqual(next(iter(openpyxl.load_workbook(out).worksheets)).freeze_panes, 'D2')

    def test_multiple_shops_become_multiple_sheets(self):
        folder = Path(self.folder.name)
        rows = folder / 'rows.jsonl'
        rows.write_text(json.dumps(['x'], ensure_ascii=False) + '\n', encoding='utf8')
        shops = [{'id': str(i), 'name': f'店{i}', 'rows': 1, 'path': str(rows), 'sheetName': f'店{i}'}
                 for i in range(1, 4)]
        index_path = folder / 'sheets.json'
        index_path.write_text(json.dumps({'headers': ['列'], 'dateColumns': [], 'numericColumns': [],
                                          'rows': 3, 'orders': 3, 'shops': shops}, ensure_ascii=False), encoding='utf8')
        out = folder / 'out.xlsx'
        self.assertEqual(xlsx_writer.write(out, index_path), 3)
        self.assertEqual(sheet_names(out), ['店1', '店2', '店3'])

    def test_progress_callback_reports_every_sheet(self):
        folder = Path(self.folder.name)
        rows = folder / 'rows.jsonl'
        rows.write_text(json.dumps(['x'], ensure_ascii=False) + '\n', encoding='utf8')
        shops = [{'id': str(i), 'name': f'店{i}', 'rows': 1, 'path': str(rows), 'sheetName': f'店{i}'}
                 for i in range(1, 4)]
        index_path = folder / 'sheets.json'
        index_path.write_text(json.dumps({'headers': ['列'], 'dateColumns': [], 'numericColumns': [],
                                          'rows': 3, 'orders': 3, 'shops': shops}, ensure_ascii=False), encoding='utf8')
        seen = []
        xlsx_writer.write(folder / 'out.xlsx', index_path, lambda done, total, name: seen.append((done, total, name)))
        self.assertEqual(seen, [(1, 3, '店1'), (2, 3, '店2'), (3, 3, '店3')], '进度要逐个 Sheet 回报（界面靠它显示进度）')


    def test_excel_limits_are_pinned(self):
        """两个上限必须是 Excel 的真实数字：写成别的值就会产出打不开的文件。"""
        self.assertEqual(xlsx_writer.MAX_ROWS, 1048576)
        self.assertEqual(xlsx_writer.MAX_COLUMNS, 16384)

    def test_too_many_columns_is_refused(self):
        """列数超限要在写出时报错，而不是产出一个 Excel 打不开的文件。"""
        original = xlsx_writer.MAX_COLUMNS
        xlsx_writer.MAX_COLUMNS = 3
        try:
            with self.assertRaisesRegex(RuntimeError, '超过 Excel 上限'):
                build(self.folder.name, '店', ['a', 'b', 'c', 'd'], [['1', '2', '3', '4']])
        finally:
            xlsx_writer.MAX_COLUMNS = original

    def test_too_many_rows_is_refused_and_leaves_no_partial_file(self):
        """行数超限同理；而且失败时不能留下半截文件（只有成功才改名成正式产物）。"""
        original = xlsx_writer.MAX_ROWS
        xlsx_writer.MAX_ROWS = 2
        try:
            with self.assertRaisesRegex(RuntimeError, '超过 Excel 上限'):
                build(self.folder.name, '店', ['列'], [['1'], ['2'], ['3']])
        finally:
            xlsx_writer.MAX_ROWS = original
        folder = Path(self.folder.name)
        # 先确认目录真的在：对不存在的目录 glob 会返回空列表且不报错，下面两条就会变成空转。
        self.assertTrue(folder.is_dir(), f'临时目录不存在，下面的检查会空转：{folder}')
        leftovers = list(folder.glob('*.part'))
        self.assertEqual(leftovers, [], f'失败后不该留下半截文件：{leftovers}')
        self.assertFalse((folder / 'out.xlsx').exists(),
                         '失败时不能产出正式产物，否则用户会拿到一个打不开的文件')


    def test_an_independent_reader_sees_the_same_values(self):
        """再用一个**独立实现**读一遍。

        此前只有 openpyxl 与项目自己的读取器读过产物，而它们与本项目同源风险较低但并非无关；
        真正的疑问是「别的电子表格程序能不能正确打开它」。python-calamine 是 Rust 实现
        （pandas 也拿它当读取引擎），与本项目代码毫无关系，是一个合适的第三方证人。
        """
        try:
            from python_calamine import CalamineWorkbook
        except ImportError:
            self.skipTest('未安装 python-calamine（开发机核对用，不在交付依赖里）')
        headers = ['订单号', '金额', '发货时间', '带空格', '特殊字符']
        rows = [['1234567890123456789', 39.9, {'date': '2026-08-31T23:59:59Z'}, '  带空格  ', 'a&b"c<d>'],
                ['000012345678901234567', 65.9, {'date': '2026-08-31T00:00:00Z'}, '普通', '中文“引号”']]
        out = build(self.folder.name, 'A"B店 & <测试>', headers, rows)
        book = CalamineWorkbook.from_path(str(out))
        self.assertEqual(book.sheet_names, ['A"B店 & <测试>'])
        table = book.get_sheet_by_name('A"B店 & <测试>').to_python()
        self.assertEqual(table[0], headers)
        first, second = table[1], table[2]
        self.assertEqual(first[0], '1234567890123456789', '19 位订单号必须是文本')
        self.assertEqual(second[0], '000012345678901234567', '前导零也必须保住——写成数字就没了')
        self.assertEqual(first[1], 39.9)
        self.assertEqual(first[3], '  带空格  ')
        self.assertEqual(first[4], 'a&b"c<d>')
        # 日期是真日期且秒不差（calamine 会把零点归成 date，两种形态都接受）
        moment = first[2]
        self.assertEqual((moment.year, moment.month, moment.day, getattr(moment, 'hour', 0),
                          getattr(moment, 'minute', 0), getattr(moment, 'second', 0)),
                         (2026, 8, 31, 23, 59, 59))


if __name__ == '__main__':
    unittest.main()
